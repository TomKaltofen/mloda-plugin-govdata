from pathlib import Path

import openpyxl
import pytest

from mloda_plugin_govdata.feature_groups.govdata.core.cache import DownloadCache
from mloda_plugin_govdata.harmonization.reference.bbsr import load_bbsr_kreise, parse_bbsr_kreise_workbook


def test_parses_all_three_fixture_sheets(fixtures_dir: Path) -> None:
    rows = parse_bbsr_kreise_workbook(fixtures_dir / "bbsr-ref-kreise-extract.xlsx")
    year_pairs = {(r.from_year, r.to_year) for r in rows}
    assert year_pairs == {(2013, 2014), (2015, 2016), (2020, 2021)}
    assert len(rows) == 21  # 8 + 7 + 6 rows across the three sheets


def test_repairs_excel_mangled_kreis_keys(fixtures_dir: Path) -> None:
    rows = parse_bbsr_kreise_workbook(fixtures_dir / "bbsr-ref-kreise-extract.xlsx")
    # Source cells store Kreise as Excel-mangled ints (e.g. 3152000); every parsed key
    # must already be the repaired 5-digit form, leading zero included.
    for row in rows:
        assert len(row.source_key) == 5
        assert len(row.target_key) == 5


def test_a_split_source_has_several_rows(fixtures_dir: Path) -> None:
    rows = parse_bbsr_kreise_workbook(fixtures_dir / "bbsr-ref-kreise-extract.xlsx")
    cochem_2013_2014 = [r for r in rows if r.from_year == 2013 and r.source_key == "07135"]
    assert len(cochem_2013_2014) == 2  # Cochem-Zell split across an identity row and a transfer row
    assert {r.target_key for r in cochem_2013_2014} == {"07135", "07140"}


def test_goettingen_merger_direction_is_old_to_new(fixtures_dir: Path) -> None:
    rows = parse_bbsr_kreise_workbook(fixtures_dir / "bbsr-ref-kreise-extract.xlsx")
    merger_rows = [r for r in rows if r.from_year == 2015 and r.source_key in {"03152", "03156"}]
    assert len(merger_rows) == 2
    assert all(r.target_key == "03159" for r in merger_rows)  # forward: old Kreise into the merged one


def test_known_upstream_defect_is_reproduced_not_fixed(fixtures_dir: Path) -> None:
    # ADR 0006 / fixture NOTICE: sheet 2015-2016 carries the 2013-2014 split shares on
    # identity rows for 07135 and 07137; this loader loads them faithfully (no share-sum
    # assertion here, that is the slice-9 re-basing loader's job).
    rows = parse_bbsr_kreise_workbook(fixtures_dir / "bbsr-ref-kreise-extract.xlsx")
    cochem_2015_2016 = next(r for r in rows if r.from_year == 2015 and r.source_key == "07135")
    assert cochem_2015_2016.target_key == "07135"  # identity row
    assert cochem_2015_2016.area_share != 1.0  # yet carries a non-identity share, the defect


def test_direction_comes_from_the_sheet_name_and_the_header_stichtage_agree(fixtures_dir: Path) -> None:
    path = fixtures_dir / "bbsr-ref-kreise-extract.xlsx"
    rows = parse_bbsr_kreise_workbook(path)
    workbook = openpyxl.load_workbook(path, read_only=True)
    for name in workbook.sheetnames:
        header = next(workbook[name].iter_rows(values_only=True))
        stichtage = tuple(int(str(header[i]).rsplit("31.12.", 1)[1]) for i in (0, 8))
        assert stichtage == tuple(int(y) for y in name.split("-"))
        assert {(r.from_year, r.to_year) for r in rows if r.from_year == stichtage[0]} == {stichtage}


SHARE_HEADERS = [
    "flächen-\nproportionaler\nUmsteige-schlüssel",
    "bevölkerungs- \nproportionaler \nUmsteige- \nschlüssel",
    "beschäftigten- \nproportionaler \nUmsteige- \nschlüssel",
    "Fläche am 31.12.2015 in km²",
    "Bevölkerung am 31.12.2015 in 1000",
    "sozialvers.pflichtig Beschäftigte am Arbeitsort am 30.6.2015 in 1000",
]


def _workbook(tmp_path: Path, header: list[str], data: list[object] | None = None) -> Path:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "2015-2016"
    sheet.append(header)
    if data is not None:
        sheet.append(data)
    path = tmp_path / "made.xlsx"
    workbook.save(path)
    return path


@pytest.mark.parametrize(
    ("first", "ninth", "seen"),
    [
        ("Kreise\n 31.12.2016", "Kreise\n 31.12.2015", r"\(2016, 2015\)"),  # swapped direction
        ("Kreise\n 30.06.2015", "Kreise\n 31.12.2016", r"\(2016,\)"),  # not the 31 Dec Stichtag
    ],
)
def test_a_header_that_contradicts_the_sheet_name_is_refused(tmp_path: Path, first: str, ninth: str, seen: str) -> None:
    path = _workbook(tmp_path, [first, "Kreisname 2015", *SHARE_HEADERS, ninth, "Kreisname 2016"])
    with pytest.raises(ValueError, match=f"sheet 2015-2016: header Stichtage {seen}"):
        parse_bbsr_kreise_workbook(path)


def test_a_header_without_share_columns_is_refused_by_name(tmp_path: Path) -> None:
    path = _workbook(tmp_path, ["Kreise\n 31.12.2015", "Kreisname 2015", "Kreise\n 31.12.2016", "Kreisname 2016"])
    with pytest.raises(ValueError, match=r"sheet 2015-2016: header lacks \['area_share', 'population_share'"):
        parse_bbsr_kreise_workbook(path)


def test_early_sheets_without_employee_columns_parse_with_none(tmp_path: Path) -> None:
    # The real file's 1990s sheets carry area and population only; columns are found by header text, not position.
    header = [
        "Kreise\n 31.12.2015",
        "Kreisname 2015",
        *SHARE_HEADERS[:2],
        *SHARE_HEADERS[3:5],
        "Kreise\n 31.12.2016",
        "Kreisname 2016",
    ]
    path = _workbook(tmp_path, header, [3152000, "Göttingen", 1, 0.5, 1117.24, 255.7, 3159000, "Göttingen"])
    (row,) = parse_bbsr_kreise_workbook(path)
    assert (row.source_key, row.area_share, row.population_share, row.target_key) == ("03152", 1.0, 0.5, "03159")
    assert (row.area_km2, row.population_thousands) == (1117.24, 255.7)
    assert (row.employee_share, row.svb_thousands) == (None, None)


@pytest.mark.live
def test_the_pinned_real_file_parses_with_the_header_cross_check(tmp_path: Path) -> None:
    with DownloadCache(tmp_path) as cache:
        rows = load_bbsr_kreise(cache, revalidate=True)
    assert {(r.from_year, r.to_year) for r in rows} == {(year, year + 1) for year in range(1990, 2024)}
    assert all(len(r.source_key) == 5 and len(r.target_key) == 5 for r in rows)
